from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documents"
DATE = "October 21, 2026"
PROJECT = "SWE Autopilot"
LONG_TITLE = (
    "AI-Driven Autonomous Software Engineering Platform with Multi-Agent "
    "Development and Intelligent Requirement Change Management"
)
ORGANIZATION = "Shaheed Zulfiqar Ali Bhutto Institute of Science and Technology University (SZABIST), Karachi Campus"
SUPERVISOR = "Adil Ahmed Unar"
TEAM = (
    "Abdul Wasay (Registration No. 2380101)\n"
    "Ahmed Hassan Sheikh (Registration No. 2380102)"
)


@dataclass(frozen=True)
class Requirement:
    req_id: str
    description: str
    priority: str = "High"
    status: str = "Pending Advisor Review"


@dataclass(frozen=True)
class Feature:
    number: str
    title: str
    description: str
    priority: str
    actors: str
    goal: str
    preconditions: str
    steps: tuple[tuple[str, str], ...]
    alternate: tuple[tuple[str, str], ...]
    exception: tuple[tuple[str, str], ...]
    postconditions: str
    requirements: tuple[Requirement, ...]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click this table of contents in Word and select Update Field."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(placeholder)
    run._r.append(fld_end)


def set_run_font(run, size=12, bold=False, italic=False, color=None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def format_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, first_line=0) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.2
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.space_after = Pt(space_after)
    fmt.first_line_indent = Inches(first_line) if first_line else None


def configure_document(doc: Document, short_title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in (("Heading 1", 16, None), ("Heading 2", 14, None), ("Heading 3", 14, None)):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        if color:
            style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.2

    if "Table Text" not in styles:
        table_style = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = styles["Table Text"]
    table_style.font.name = "Times New Roman"
    table_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    table_style.font.size = Pt(10.5)
    table_style.paragraph_format.line_spacing = 1.05
    table_style.paragraph_format.space_after = Pt(0)

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run, size=10, italic=True)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        set_run_font(run, size=10)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    format_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix) :])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_bullets(doc: Document, items: Iterable[str], level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=True)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None, header_fill="1F4E78"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[index])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.style = "Table Text"
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color=(255, 255, 255))
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[index])
            p = cell.paragraphs[0]
            p.style = "Table Text"
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=10.5)
    return table


def add_cover(doc: Document, document_type: str, subtitle: str) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(document_type)
    set_run_font(r, size=20, bold=True)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("for")
    set_run_font(r, size=14, bold=True)
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PROJECT)
    set_run_font(r, size=18, bold=True)
    p.paragraph_format.space_after = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(LONG_TITLE)
    set_run_font(r, size=13, bold=True)
    p.paragraph_format.space_after = Pt(20)

    for label, value in (
        ("Document status", "Draft for Advisor Review"),
        ("Scope", "FYP-I / Version 1.0"),
        ("Prepared by", TEAM),
        ("Organization", ORGANIZATION),
        ("Supervisor", SUPERVISOR),
        ("Date created", DATE),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        first = p.add_run(f"{label}: ")
        set_run_font(first, size=12, bold=True)
        second = p.add_run(value)
        set_run_font(second, size=12)
    doc.add_page_break()


def add_front_matter(doc: Document, include_revision=True) -> None:
    add_heading(doc, "Table of Contents", 1)
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
    add_toc(p)
    doc.add_page_break()
    if include_revision:
        add_heading(doc, "Revision History", 1)
        add_table(
            doc,
            ["Name", "Date", "Reason For Changes", "Version"],
            [["Abdul Wasay / Ahmed Hassan Sheikh", DATE, "Initial FYP-I SRS and use-case baseline for advisor review", "1.0"]],
            widths=[1.7, 1.2, 3.6, 0.8],
            header_fill="1F4E78",
        )
        doc.add_page_break()


def add_use_case_tables(doc: Document, feature: Feature, table_index: int, include_title=True) -> int:
    if include_title:
        add_heading(doc, f"4.{feature.number}. {feature.title}", 2)
        p = doc.add_paragraph()
        format_paragraph(p)
        r = p.add_run("Description: ")
        set_run_font(r, bold=True)
        r = p.add_run(feature.description)
        set_run_font(r)

    add_heading(doc, f"4.{feature.number}.1. Description and Priority", 3)
    add_body(doc, f"{feature.description} Priority: {feature.priority}.")

    add_heading(doc, f"4.{feature.number}.2. Stimulus/Response Sequences", 3)
    summary_rows = [
        ["Use Case Name", feature.title],
        ["Description", feature.description],
        ["Actors", feature.actors],
        ["Goal", feature.goal],
        ["Pre-Conditions", feature.preconditions],
    ]
    add_table(doc, ["Use Case Field", "Specification"], summary_rows, widths=[1.7, 5.6], header_fill="1F4E78")
    add_caption(doc, f"Table {table_index}. {feature.title} Use Case Summary")
    table_index += 1

    step_rows = []
    for index, (action, response) in enumerate(feature.steps, start=1):
        step_rows.append([str(index), action, response])
    add_table(doc, ["Step", "User/Agent Action", "System Response"], step_rows, widths=[0.6, 3.0, 3.7], header_fill="1F4E78")
    add_caption(doc, f"Table {table_index}. {feature.title} Basic Course of Events")
    table_index += 1

    add_heading(doc, f"4.{feature.number}.2.1. Alternate Flow", 3)
    alt_rows = [[action, response] for action, response in feature.alternate]
    add_table(doc, ["User/Agent Action", "System Response"], alt_rows, widths=[3.0, 4.3], header_fill="5B7FA3")
    add_caption(doc, f"Table {table_index}. {feature.title} Alternate Flow")
    table_index += 1

    add_heading(doc, f"4.{feature.number}.2.2. Exception Flow", 3)
    ex_rows = [[action, response] for action, response in feature.exception]
    add_table(doc, ["Exception Condition", "System Response"], ex_rows, widths=[3.0, 4.3], header_fill="5B7FA3")
    add_caption(doc, f"Table {table_index}. {feature.title} Exception Flow")
    table_index += 1

    p = doc.add_paragraph()
    format_paragraph(p)
    r = p.add_run("Post Conditions: ")
    set_run_font(r, bold=True)
    r = p.add_run(feature.postconditions)
    set_run_font(r)

    add_heading(doc, f"4.{feature.number}.3. Functional Requirements", 3)
    req_rows = [[r.req_id, r.description, r.priority, r.status] for r in feature.requirements]
    add_table(doc, ["Req.ID", "Description", "Priority", "Status"], req_rows, widths=[0.85, 4.4, 0.9, 1.15], header_fill="1F4E78")
    add_caption(doc, f"Table {table_index}. {feature.title} Functional Requirements")
    table_index += 1
    return table_index


def add_feature_srs(doc: Document, feature: Feature) -> None:
    add_heading(doc, f"4.{feature.number} {feature.title}", 2)
    add_heading(doc, f"4.{feature.number}.1 Feature Summary", 3)
    summary_rows = [
        ["Feature ID", f"FYP1-F{feature.number}"],
        ["Feature Name", feature.title],
        ["Description", feature.description],
        ["Priority", feature.priority],
        ["Scope", "FYP-I baseline; detailed flows are documented in the separate System Use Case Specification."],
    ]
    add_table(doc, ["Feature Field", "Specification"], summary_rows, widths=[1.7, 5.6], header_fill="1F4E78")
    add_caption(doc, f"Table 4.{feature.number}.1. {feature.title} Feature Summary")
    add_heading(doc, f"4.{feature.number}.2 Functional Requirements", 3)
    req_rows = [[r.req_id, r.description, r.priority, r.status] for r in feature.requirements]
    add_table(doc, ["Req.ID", "Description", "Priority", "Status"], req_rows, widths=[0.85, 4.4, 0.9, 1.15], header_fill="1F4E78")
    add_caption(doc, f"Table 4.{feature.number}.2. {feature.title} Functional Requirements")


def add_references(doc: Document) -> None:
    refs = [
        "SWE Autopilot FYP Proposal, Complete Unified Document, submitted October 7, 2026, SZABIST Karachi Campus.",
        "Software Requirements Specification Template (2025), Karl E. Wiegers template, supplied by the project advisor.",
        "Sign Up Use Case Example, supplied by the project advisor for Section 4 formatting guidance.",
        "SWE Autopilot implementation repository: https://github.com/sheikh4467/SWE-Autopilot",
        "FastAPI documentation: https://fastapi.tiangolo.com/",
        "Next.js documentation: https://nextjs.org/docs",
        "LangGraph documentation: https://langchain-ai.github.io/langgraph/",
        "Docker documentation: https://docs.docker.com/",
    ]
    add_numbered(doc, refs)


def add_srs_document(features: Sequence[Feature]) -> Path:
    doc = Document()
    configure_document(doc, f"Software Requirements Specification for {PROJECT}")
    add_cover(doc, "Software Requirements Specification", "FYP-I")
    add_front_matter(doc)

    add_heading(doc, "1. Introduction", 1)
    add_heading(doc, "1.1 Purpose", 2)
    add_body(doc, "This Software Requirements Specification (SRS) defines the functional and nonfunctional requirements for the FYP-I release of SWE Autopilot. SWE Autopilot is an AI-driven autonomous software engineering platform that coordinates specialized agents across requirements analysis, architecture, code generation, execution, and verification. This document specifies the product baseline that will be reviewed by the project advisor before implementation and submission. Detailed use-case flows and the ERD are delivered as separate documents.")
    add_body(doc, "The FYP-I scope is limited to the requirement analysis and ambiguity detector, requirement-to-artifact traceability, predictive risk analysis, Agent Runtime foundations, Architect/Critic review, multi-agent code generation, Docker sandbox execution, initial evidence-based verification, project workspace, and supporting authentication. FYP-II features such as the What-If Simulator, full change-impact engine, closed-loop repair, Attack Memory, rollback, deployment governance, and final SDLC reporting are recorded as future scope and are not treated as FYP-I acceptance requirements.")

    add_heading(doc, "1.2 Document Conventions", 2)
    add_body(doc, "This document uses Times New Roman. Body text is 12-point with 1.2 line spacing, normal margins, and justified paragraphs. Main headings use 16-point bold text and subsection headings use 14-point bold text. Functional requirements use unique identifiers. Priority values are High, Medium, or Low. Status values are marked Pending Advisor Review until the supervisor approves the baseline. The terms shall and must identify mandatory requirements; should identifies a preferred requirement; and may identifies an optional capability.")

    add_heading(doc, "1.3 Intended Audience and Reading Suggestions", 2)
    add_body(doc, "The primary audience is the FYP advisor, project team, software developers, testers, and reviewers. Advisors should review Sections 1, 2, 4, and 5 to confirm scope and acceptance criteria. Developers should use Sections 3 and 4 for interfaces and requirements. Testers should use Section 4 functional requirements and Section 5 quality requirements to derive test cases. Detailed stimulus/response flows are provided in the separate System Use Case Specification, while the database model is provided in the separate ERD document.")

    add_heading(doc, "1.4 Product Scope", 2)
    add_body(doc, "SWE Autopilot provides a web-based control plane for turning natural-language software requirements into structured, traceable, and verifiable engineering work. The platform places a deterministic Agent Runtime around LLM-based reasoning. It controls workflow state, agent permissions, artifact versions, execution evidence, human approval points, and project history rather than treating generated code as automatically correct. The FYP-I product demonstrates the foundation of this workflow for a controlled full-stack web application scenario.")
    add_bullets(doc, [
        "Analyze natural-language requirements and identify ambiguity, incompleteness, conflicts, and missing acceptance criteria.",
        "Create requirement IDs, user stories, acceptance criteria, and artifact relationships.",
        "Predict architecture, technical, performance, security, dependency, and implementation risks.",
        "Generate architecture proposals and critique them before implementation.",
        "Generate application code through specialized agents under controlled permissions.",
        "Execute generated software in a Docker sandbox and collect build and runtime evidence.",
        "Run the initial verification pipeline and associate results with project artifacts.",
    ])

    add_heading(doc, "1.5 References", 2)
    add_references(doc)

    add_heading(doc, "2. Overall Description", 1)
    add_heading(doc, "2.1 Product Perspective", 2)
    add_body(doc, "SWE Autopilot is a new, self-contained web platform composed of a Next.js frontend, FastAPI backend, Agent Runtime, specialized SDLC agents, data storage, and controlled execution services. The frontend provides the landing/authentication experience, project workspace, requirement input, artifact views, architecture review, and verification status. The backend exposes REST services, coordinates the workflow, stores artifacts and audit information, and calls configured LLM providers. Docker provides the execution boundary for generated projects.")
    add_body(doc, "The target FYP architecture uses PostgreSQL for project data, pgvector for vector storage where required, Redis for task and memory support, LangGraph/LangChain for orchestration, and Docker for execution. The current authentication prototype uses SQLite locally; this is an implementation convenience and does not change the target platform architecture.")

    add_heading(doc, "2.2 Product Functions", 2)
    add_bullets(doc, [
        "Register, sign in, and sign out using email/password and configured Google/GitHub OAuth providers.",
        "Create a project workspace and submit a natural-language software requirement.",
        "Compile requirements into structured artifacts and flag ambiguity or incomplete acceptance criteria.",
        "Maintain a traceability matrix connecting requirements to user stories, architecture, APIs, code, tests, documentation, and evidence.",
        "Analyze risks and present severity, rationale, affected artifacts, and mitigation suggestions.",
        "Generate architecture alternatives and coordinate an Architect/Critic review with an approval gate.",
        "Generate application artifacts through specialized agents with explicit repository and tool permissions.",
        "Build and run generated projects inside restricted Docker environments.",
        "Collect build, startup, test, static-analysis, schema, and runtime evidence for the initial verification result.",
        "Record artifact versions, execution history, approval decisions, and unresolved issues for the project workspace.",
    ])

    add_heading(doc, "2.3 User Classes and Characteristics", 2)
    add_table(doc, ["User Class", "Characteristics and Responsibilities"], [
        ["Project Member / Developer", "Creates workspaces, enters requirements, reviews artifacts, approves architecture, initiates generation, and inspects verification evidence. Has technical software-engineering knowledge."],
        ["Advisor / Reviewer", "Reviews requirements, architecture decisions, evidence, scope, and unresolved issues. May approve or reject an FYP-I milestone."],
        ["System Administrator", "Configures model providers, repository access, execution limits, and environment settings. This role is limited in the FYP-I prototype."],
        ["Specialized Agent / Runtime", "Internal system actors that compile requirements, analyze risks, propose or critique architecture, generate code, and coordinate execution. Agents do not independently bypass runtime permissions."],
    ], widths=[2.0, 5.3])

    add_heading(doc, "2.4 Operating Environment", 2)
    add_body(doc, "The platform shall operate as a browser-accessible web application. Development and demonstration environments may use Windows, Linux, or macOS with a multi-core processor and at least 16 GB RAM recommended. The target stack includes Python 3.11 or later, Node.js, Docker Desktop, PostgreSQL, Redis, Git, a supported modern web browser, and access to an LLM API or local Ollama model. The frontend is implemented with Next.js, React, TypeScript, Tailwind CSS, and custom CSS motion effects. The backend is implemented with Python and FastAPI.")

    add_heading(doc, "2.5 Design and Implementation Constraints", 2)
    add_bullets(doc, [
        "The system shall preserve a deterministic workflow around LLM calls; an LLM response alone shall not mark an artifact as verified.",
        "Generated code shall be executed only through a restricted Docker sandbox and shall not receive unrestricted host access.",
        "The Agent Runtime shall control tool permissions, repository access, artifact versions, state transitions, and approval gates.",
        "OAuth client secrets, session secrets, and LLM API keys shall be loaded from protected environment variables and shall not be stored in frontend code or source control.",
        "The initial platform targets full-stack web applications; native desktop, embedded, hardware, and native mobile compilers are outside FYP-I scope.",
        "The architecture shall remain model-provider independent so that OpenAI, Gemini, or Ollama can be selected without rewriting the workflow core.",
    ])

    add_heading(doc, "2.6 User Documentation", 2)
    add_bullets(doc, [
        "Quick-start guide for creating a workspace and entering requirements.",
        "User guide for reviewing ambiguity findings, traceability links, risks, architecture proposals, and verification evidence.",
        "Developer setup guide for frontend, backend, OAuth, Docker, and model-provider configuration.",
        "Advisor review checklist covering FYP-I feature acceptance, evidence, and unresolved issues.",
    ])

    add_heading(doc, "2.7 Assumptions and Dependencies", 2)
    add_bullets(doc, [
        "Users have access to a modern browser and a reachable SWE Autopilot deployment.",
        "The project has access to Docker and the required runtime images for sandbox execution.",
        "At least one configured LLM provider or local model is available for agent reasoning and generation.",
        "The target repository is a Git repository that can be read and modified under explicit permission controls.",
        "Google/GitHub OAuth credentials are correctly configured when social authentication is enabled.",
        "PostgreSQL and Redis are available for the target deployment; the local authentication prototype may use SQLite during development.",
    ])

    add_heading(doc, "3. External Interface Requirements", 1)
    add_heading(doc, "3.1 User Interfaces", 2)
    add_body(doc, "The frontend shall provide a responsive landing/authentication page that explains who SWE Autopilot is and what it does, with 3D orbital agent-runtime visuals, scroll transitions, and clear Sign In and Create Account actions. The authenticated workspace shall provide navigation to project requirements, traceability, risk analysis, architecture review, code generation, sandbox execution, and verification evidence. Forms shall show field-level validation and actionable error messages. Tables and status badges shall show requirement IDs, priorities, artifact relationships, approval status, execution status, and verification evidence.")
    add_heading(doc, "3.2 Hardware Interfaces", 2)
    add_body(doc, "No dedicated hardware interface is required. The system shall use standard computer hardware through the operating system. Docker execution may use CPU, memory, storage, and network resources subject to configured sandbox limits. The platform shall not require access to sensors, embedded boards, or hardware controllers in FYP-I.")
    add_heading(doc, "3.3 Software Interfaces", 2)
    add_table(doc, ["Interface", "Purpose and Data Exchanged"], [
        ["Next.js frontend ↔ FastAPI backend", "REST/JSON requests for authentication, projects, requirements, artifacts, agent execution, and verification status."],
        ["FastAPI ↔ PostgreSQL", "Project, requirement, artifact, relationship, risk, architecture, execution, and evidence records."],
        ["FastAPI ↔ Redis", "Task state, queues, temporary workflow data, and memory support where enabled."],
        ["Agent Runtime ↔ LangGraph/LangChain", "State transitions, agent messages, tool requests, permissions, and execution outcomes."],
        ["Agent Runtime ↔ LLM provider", "Structured prompts and model responses for analysis, architecture, generation, and critique."],
        ["Runtime ↔ Docker", "Build/run commands, mounted project artifacts, resource restrictions, logs, exit status, and test results."],
        ["Authentication ↔ Google/GitHub OAuth", "Authorization code, provider identity, email, and profile data through server-side callbacks."],
    ], widths=[2.0, 5.3])
    add_heading(doc, "3.4 Communications Interfaces", 2)
    add_body(doc, "Browser-to-backend communication shall use HTTP during local development and HTTPS in deployed environments. REST requests shall use JSON except for browser redirects required by OAuth. OAuth state shall be validated server-side. Authenticated sessions shall use HTTP-only cookies. LLM and provider communications shall use TLS-protected endpoints. Docker logs and verification results shall be transferred to the backend through controlled local process or service interfaces.")

    add_heading(doc, "4. System Features", 1)
    add_body(doc, "Section 4 defines the FYP-I system features and their functional requirements at SRS level. Detailed actors, preconditions, stimulus/response sequences, alternate flows, exception flows, and postconditions are maintained in the separate SWE Autopilot FYP-I System Use Case Specification. All statuses remain Pending Advisor Review until the supervisor approves the baseline.")
    for feature in features:
        add_feature_srs(doc, feature)

    add_heading(doc, "5. Other Nonfunctional Requirements", 1)
    add_heading(doc, "5.1 Performance Requirements", 2)
    add_bullets(doc, [
        "The landing page shall become interactive within 3 seconds on the target local development environment after assets are available.",
        "The backend shall return validation errors for invalid form input within 2 seconds under normal local load.",
        "A requirement analysis request shall show progress or a running state within 3 seconds of submission, even if the LLM analysis takes longer.",
        "The workspace shall display the latest saved artifact and status records within 3 seconds for a project containing up to 1,000 FYP-I artifacts.",
        "Docker execution shall stream or persist logs so that a user can inspect progress without waiting for the entire process to finish.",
    ])
    add_heading(doc, "5.2 Safety Requirements", 2)
    add_bullets(doc, [
        "The system shall not execute generated code directly on the host operating system when a sandbox is available.",
        "The sandbox shall enforce time, memory, storage, and process limits where supported.",
        "The runtime shall stop an execution that exceeds its configured limit and shall record the reason.",
        "The system shall require explicit approval before an architecture decision or generated artifact is promoted to the next controlled stage.",
    ])
    add_heading(doc, "5.3 Security Requirements", 2)
    add_bullets(doc, [
        "The system shall support email/password authentication and configured Google and GitHub OAuth authentication.",
        "Passwords shall be stored only as strong one-way hashes; plaintext passwords shall never be persisted.",
        "OAuth state shall be generated, stored, validated, and consumed to prevent callback tampering.",
        "Authenticated sessions shall use HTTP-only cookies and shall not expose provider secrets or access tokens in URLs.",
        "Secrets including OAuth client secrets, LLM keys, and session keys shall be supplied through protected environment variables.",
        "The runtime shall restrict tools and artifacts available to each agent and shall retain an audit record of sensitive actions.",
        "Project data shall be isolated by workspace and user authorization checks.",
    ])
    add_heading(doc, "5.4 Software Quality Attributes", 2)
    add_table(doc, ["Attribute", "Requirement"], [
        ["Usability", "A first-time project member should be able to create a workspace, submit a requirement, and locate the first analysis result without external assistance."],
        ["Correctness", "A feature shall not be marked verified without recorded execution or validation evidence appropriate to that feature."],
        ["Reliability", "A failed agent or Docker execution shall produce a visible failure state and diagnostic record rather than silently losing workflow state."],
        ["Maintainability", "Frontend, backend, runtime, agents, and storage adapters shall be separated into maintainable modules."],
        ["Portability", "The platform shall run in local Windows/Linux/macOS development environments and in a container-capable deployment environment."],
        ["Testability", "Functional requirements shall have identifiable inputs, outputs, error paths, and evidence suitable for automated tests."],
        ["Traceability", "Requirements and generated artifacts shall carry stable IDs and relationships that can be queried and reviewed."],
    ], widths=[1.5, 5.8])
    add_heading(doc, "5.5 Business Rules", 2)
    add_bullets(doc, [
        "BR-01: A project member may initiate an analysis only for a workspace to which the member has access.",
        "BR-02: An artifact may not be marked verified solely because an LLM generated or recommended it.",
        "BR-03: An architecture decision requires a recorded rationale and an approval decision before code generation is promoted.",
        "BR-04: Generated code must remain linked to the requirement and architecture artifacts that motivated it.",
        "BR-05: A failed or incomplete verification result shall remain visible to the user and shall not be silently converted to success.",
        "BR-06: FYP-II capabilities shall not be represented as completed FYP-I functionality in the advisor review baseline.",
    ])

    add_heading(doc, "6. Other Requirements", 1)
    add_bullets(doc, [
        "The system shall maintain timestamps and version identifiers for requirements, artifacts, architecture records, execution runs, and evidence.",
        "The system shall support export of a project summary containing requirements, architecture decisions, risk findings, execution evidence, and unresolved issues.",
        "The system shall use UTF-8 text handling for natural-language requirements and generated documentation.",
        "The system shall keep FYP-I and FYP-II scope labels visible in project planning and advisor review views.",
    ])

    add_heading(doc, "Appendix A: Glossary", 1)
    add_table(doc, ["Term", "Definition"], [
        ["Agent Runtime", "The deterministic control layer that manages workflow state, permissions, artifacts, evidence, approvals, and audit history."],
        ["Architect Agent", "An agent that proposes components, APIs, data models, technologies, and communication patterns."],
        ["Critic Agent", "An agent that challenges an architecture proposal for security, scalability, complexity, dependency, and requirement issues."],
        ["Artifact", "A managed project object such as a requirement, user story, architecture record, API, source file, test, documentation item, or evidence record."],
        ["Traceability", "The set of relationships connecting requirements to downstream software artifacts and verification evidence."],
        ["Ambiguity", "A requirement condition where multiple interpretations are possible or important information is missing."],
        ["Docker Sandbox", "A restricted execution environment used to build and run generated software without unrestricted host access."],
        ["Verification Evidence", "Observable build, runtime, test, static-analysis, schema, or execution output used to support a result."],
        ["FYP-I", "The first semester deliverable focused on requirements, runtime foundations, architecture, code generation, sandboxing, and initial verification."],
        ["LLM", "Large Language Model used for reasoning or generation inside the controlled workflow."],
    ], widths=[2.0, 5.3])

    add_heading(doc, "Appendix B: Analysis Models", 1)
    add_body(doc, "Top-level FYP-I data flow:")
    add_body(doc, "Project Member → Requirement Input → Requirement Compiler → Structured Requirements / User Stories / Acceptance Criteria → Traceability Store → Risk Analyzer → Architect/Critic → Approved Architecture → Code Agents → Docker Sandbox → Verification Evidence → Project Workspace.")
    add_body(doc, "Primary traceability path:")
    add_body(doc, "REQ-### → US-### → ARCH-### → API/COMPONENT → CODE ARTIFACT → TEST ARTIFACT → EXECUTION RUN → VERIFICATION EVIDENCE.")
    add_body(doc, "FYP-I state flow:")
    add_body(doc, "Draft Requirement → Analyzed → Reviewed → Architecture Proposed → Architecture Approved → Code Generated → Sandbox Executed → Verification Passed / Verification Failed / Escalated.")
    add_body(doc, "The complete logical ERD is maintained as the separate document SWE_Autopilot_Whole_Project_ERD_CrowFoot.pdf and is intentionally not embedded in this SRS.")

    add_heading(doc, "Appendix C: To Be Determined List", 1)
    add_table(doc, ["TBD ID", "Open Item", "Owner / Resolution Point"], [
        ["TBD-01", "Final PostgreSQL schema and migration strategy for all FYP-I artifacts.", "Team; resolve during Agent Runtime and platform-core design."],
        ["TBD-02", "Exact LLM model/provider selected for the advisor demonstration.", "Team and advisor; resolve before code-generation evaluation."],
        ["TBD-03", "Final Docker resource limits and permitted network policy.", "Team; resolve before sandbox verification."],
        ["TBD-04", "Detailed advisor approval workflow and roles in the final UI.", "Advisor and team; resolve during review-gate design."],
        ["TBD-05", "Benchmark dataset and acceptance thresholds for ambiguity and risk analysis.", "Team; resolve before FYP-I evaluation."],
    ], widths=[1.0, 4.2, 2.1])

    path = OUT / "SWE_Autopilot_FYP1_SRS.docx"
    doc.save(path)
    return path


def add_use_case_document(features: Sequence[Feature]) -> Path:
    doc = Document()
    configure_document(doc, f"FYP-I System Use Cases for {PROJECT}")
    add_cover(doc, "FYP-I System Use Case Specification", "FYP-I")
    add_front_matter(doc)
    add_heading(doc, "Purpose of This Document", 1)
    add_body(doc, "This document expands Section 4 of the SWE Autopilot Software Requirements Specification into advisor-reviewable system use cases. It follows the supplied Sign Up Use Case Example: every use case includes a description, actors, goal, preconditions, basic course of events, alternate flow, exception flow, postconditions, a stimulus/response table, and a functional requirements table.")
    add_body(doc, "The use cases are limited to the FYP-I feature set from the project proposal and the supporting authentication/workspace flow implemented for the current prototype. All requirement statuses are Pending Advisor Review and must be checked by the advisor before submission.")
    table_index = 1
    for index, feature in enumerate(features):
        if index > 0:
            doc.add_page_break()
        table_index = add_use_case_tables(doc, feature, table_index, include_title=True)
    add_heading(doc, "FYP-I Scope Boundary", 1)
    add_body(doc, "The following proposal features are intentionally excluded from these FYP-I use cases and remain FYP-II scope: What-If Development Simulator, Change-Impact and Consistency Engine, Testing Agent and selective test replay, Closed-Loop Debugging and Recovery, Defensive Red-Team Agent, Attack Memory, Controlled Rollback, Release Gate, Deployment, final SDLC reporting, benchmarking, and final defense preparation.")
    add_heading(doc, "Advisor Review Checklist", 1)
    add_bullets(doc, [
        "Confirm that each FYP-I feature has a clear actor, goal, precondition, main flow, alternate flow, exception flow, and postcondition.",
        "Confirm that functional requirements are testable, necessary, unambiguous, and correctly prioritized.",
        "Confirm that the authentication and workspace requirements are treated as supporting platform functions, not as a replacement for the core FYP-I research contribution.",
        "Confirm which requirements are accepted, revised, deferred, or removed before implementation and final submission.",
        "Update each Status column after advisor review and record the revision in the Revision History table.",
    ])
    add_heading(doc, "References", 1)
    add_references(doc)

    path = OUT / "SWE_Autopilot_FYP1_System_Use_Cases.docx"
    doc.save(path)
    return path


FEATURES: tuple[Feature, ...] = (
    Feature(
        "1", "User Registration and Authentication",
        "The platform allows a project member to create an account and access the SWE Autopilot workspace using email/password authentication or configured Google and GitHub OAuth.",
        "High", "Project Member; Authentication Service; Google OAuth; GitHub OAuth", "Create or access a protected project workspace.", "The user is not authenticated; the backend and configured provider credentials are available for the selected method.",
        (
            ("The user opens the landing page and selects Sign In or Get Started.", "The system displays the authentication modal with email/password fields and Google/GitHub options."),
            ("The user enters a name, email, and password for sign-up, or an email and password for sign-in.", "The system validates required fields, email format, and password length before submission."),
            ("The user submits valid email/password data.", "The backend checks the account, hashes a new password for sign-up or verifies the stored hash for sign-in, creates an HTTP-only session, and returns the public user record."),
            ("The user selects Google or GitHub.", "The system redirects the browser to the backend OAuth route, validates the callback state, links or creates the user, establishes a session, and returns the user to the frontend."),
            ("The authenticated user selects Sign Out.", "The system clears the session and returns the user to the unauthenticated landing state."),
        ),
        (("The user submits blank or invalid fields.", "The system highlights the invalid input and prevents submission until the data is corrected."), ("The email already exists during sign-up.", "The system displays an account-exists message and suggests signing in.")),
        (("Database, provider, or network service is unavailable.", "The system displays a recoverable error, does not expose secrets, and keeps the user on the authentication screen."), ("OAuth state is missing, invalid, or expired.", "The backend rejects the callback and displays an authentication failure message.")),
        "An authenticated user has a valid session and can create or open a project workspace; after sign-out, the session is cleared.",
        (
            Requirement("AUTH-01", "The system shall display sign-up and sign-in forms with required name, email, and password fields as appropriate."),
            Requirement("AUTH-02", "The system shall validate email format, required fields, and a minimum password length of eight characters before submission."),
            Requirement("AUTH-03", "The system shall prevent duplicate local accounts for the same normalized email address."),
            Requirement("AUTH-04", "The system shall hash passwords using a one-way password hashing function and shall never store plaintext passwords."),
            Requirement("AUTH-05", "The system shall support configured Google and GitHub OAuth login through server-side callback routes."),
            Requirement("AUTH-06", "The system shall create an HTTP-only authenticated session containing only the application user identifier."),
            Requirement("AUTH-07", "The system shall display actionable errors for invalid credentials, provider failures, and unavailable services."),
        ),
    ),
    Feature(
        "2", "Requirement Analysis and Ambiguity Detection",
        "The Requirement Analyzer converts a natural-language software requirement into structured requirement artifacts and identifies ambiguity, incompleteness, conflicts, and missing acceptance criteria.",
        "High", "Project Member; Requirement Compiler Agent; Agent Runtime", "Produce a reviewed, structured requirement baseline for a project.", "The user is authenticated, a project workspace exists, and the requirement text is available.",
        (
            ("The user opens a project workspace and enters a natural-language requirement.", "The system displays the requirement editor and accepts the text for analysis."),
            ("The user selects Analyze Requirement.", "The Agent Runtime creates an analysis run, validates the input, and invokes the Requirement Compiler Agent."),
            ("The compiler processes the requirement.", "The system extracts functional requirements, nonfunctional requirements, user stories, acceptance criteria, IDs, entities, and likely dependencies."),
            ("The analysis completes.", "The system highlights ambiguous statements, missing information, conflicts, and incomplete acceptance criteria with explanations and severity."),
            ("The user reviews or edits the findings.", "The system stores the accepted requirement baseline and records the review state and version."),
        ),
        (("The requirement is empty or too short.", "The system requests meaningful requirement text and does not start an LLM analysis."), ("The user rejects a generated interpretation.", "The system preserves the original input, marks the interpretation as rejected, and allows revision and re-analysis.")),
        (("The model provider is unavailable or returns an invalid response.", "The runtime marks the run as failed, records the diagnostic, and does not create unvalidated requirement artifacts."), ("The response violates the expected schema.", "The system rejects the response and requests a bounded retry or human review.")),
        "The project contains a versioned requirement analysis with structured artifacts, identified ambiguities, and a visible review status.",
        (
            Requirement("REQ-AN-01", "The system shall accept a natural-language requirement for a selected project."),
            Requirement("REQ-AN-02", "The system shall create unique IDs for extracted requirements, user stories, and acceptance criteria."),
            Requirement("REQ-AN-03", "The system shall classify extracted items as functional or nonfunctional requirements where applicable."),
            Requirement("REQ-AN-04", "The system shall identify ambiguous, incomplete, conflicting, or underspecified requirement statements."),
            Requirement("REQ-AN-05", "The system shall explain each ambiguity or missing information finding in user-readable language."),
            Requirement("REQ-AN-06", "The system shall allow the user to review, edit, accept, or reject generated requirement artifacts."),
            Requirement("REQ-AN-07", "The system shall preserve the original requirement text and the version history of analysis results."),
        ),
    ),
    Feature(
        "3", "Requirement-to-Artifact Traceability",
        "The Traceability Manager links requirements to user stories, architecture components, APIs, code artifacts, tests, documentation, and verification evidence.",
        "High", "Project Member; Traceability Service; Agent Runtime", "Make the relationship between a requirement and downstream engineering evidence visible and reviewable.", "A project contains at least one analyzed requirement and one or more generated or manually linked artifacts.",
        (
            ("The user opens the Traceability view.", "The system loads the project graph and displays requirement IDs and linked artifacts."),
            ("The system compiles or updates artifacts.", "The Traceability Service creates or updates links such as REQ → US → ARCH → API/COMPONENT → CODE → TEST → EVIDENCE."),
            ("The user selects a requirement or artifact.", "The system shows direct links, dependent links, artifact versions, and missing downstream links."),
            ("The user adds or corrects a relationship.", "The system validates the relationship, saves it with the current project version, and records the change."),
        ),
        (("An artifact cannot be uniquely identified.", "The system requests a valid artifact ID or marks the relationship as unresolved."), ("A downstream artifact is missing.", "The system shows a traceability gap and identifies the next artifact expected by the workflow.")),
        (("The graph store or database is unavailable.", "The system displays the last saved traceability snapshot and reports that new changes could not be saved."), ("A circular or invalid relationship is submitted.", "The system rejects the relationship and explains the allowed relationship types.")),
        "The project has a queryable, versioned traceability matrix with visible gaps and evidence links.",
        (
            Requirement("TRACE-01", "The system shall display requirement-to-artifact relationships for the selected project."),
            Requirement("TRACE-02", "The system shall support relationships between requirements, user stories, architecture, APIs/components, code, tests, documentation, and evidence."),
            Requirement("TRACE-03", "The system shall preserve stable IDs and version references for linked artifacts."),
            Requirement("TRACE-04", "The system shall identify missing or unresolved links in the traceability chain."),
            Requirement("TRACE-05", "The system shall allow an authorized user to add, edit, or remove a relationship."),
            Requirement("TRACE-06", "The system shall record relationship changes in the project audit history."),
        ),
    ),
    Feature(
        "4", "Predictive Risk Analysis",
        "The Predictive Risk Analyzer evaluates requirements and proposed architecture for technical, architectural, performance, security, dependency, ambiguity, and implementation-complexity risks.",
        "High", "Project Member; Risk Analyzer Agent; Agent Runtime", "Identify important risks before implementation begins and provide mitigation guidance.", "A reviewed requirement baseline and, where available, an architecture proposal exist.",
        (
            ("The user selects Run Risk Analysis.", "The runtime creates a risk-analysis run and collects the relevant requirements, artifacts, dependencies, and architecture records."),
            ("The Risk Analyzer evaluates the project.", "The system identifies risk factors, affected artifacts, likelihood, impact, rationale, and suggested mitigation."),
            ("The analysis completes.", "The system groups risks by category and sorts them by severity or priority."),
            ("The user reviews a risk item.", "The system shows the evidence or artifacts that led to the risk and allows the user to acknowledge, annotate, or mark it for follow-up."),
        ),
        (("The user runs analysis without a requirement baseline.", "The system requests requirement analysis first and does not produce an ungrounded risk report."), ("The user marks a risk as accepted.", "The system records the acceptance rationale and keeps the risk visible in project history.")),
        (("The model provider or dependency graph is unavailable.", "The system records a partial or failed analysis and identifies the missing input instead of claiming that no risks exist."), ("A risk response is malformed.", "The runtime rejects the response, records the schema error, and requests review or a bounded retry.")),
        "The project contains a categorized, prioritized risk register with rationale, affected artifacts, mitigation suggestions, and review status.",
        (
            Requirement("RISK-01", "The system shall run a risk analysis for a selected project baseline."),
            Requirement("RISK-02", "The system shall evaluate at least architecture, technical, performance, security, dependency, ambiguity, and implementation-complexity risk categories."),
            Requirement("RISK-03", "The system shall assign a severity or priority to each identified risk."),
            Requirement("RISK-04", "The system shall identify the requirement or artifact affected by each risk where possible."),
            Requirement("RISK-05", "The system shall provide a rationale and mitigation suggestion for each risk."),
            Requirement("RISK-06", "The system shall allow an authorized user to acknowledge, annotate, or defer a risk with a recorded rationale."),
        ),
    ),
    Feature(
        "5", "Architect/Critic Review and Approval",
        "The Architect Agent proposes a system architecture and the Critic Agent evaluates it for security, scalability, complexity, missing components, dependency problems, and requirement mismatches.",
        "High", "Project Member; Architect Agent; Critic Agent; Agent Runtime; Advisor/Reviewer", "Produce an approved architecture decision with alternatives, critique, rationale, and an auditable approval state.", "A structured requirement baseline and relevant risk analysis are available.",
        (
            ("The user requests architecture proposals.", "The runtime provides the approved requirements and risk findings to the Architect Agent."),
            ("The Architect Agent creates one or more alternatives.", "The system displays components, APIs, data design, technology choices, communication patterns, assumptions, and trade-offs."),
            ("The Critic Agent reviews the proposal.", "The system displays findings related to security, scalability, complexity, dependencies, missing components, and requirement coverage."),
            ("The user or advisor reviews the debate.", "The runtime records discussion messages, revisions, and the selected architecture decision."),
            ("An authorized reviewer approves or rejects the decision.", "The system records the approval result and allows code generation only for an approved architecture."),
        ),
        (("The critic identifies blocking issues.", "The system returns the proposal to the architecture-review state and requires revision before approval."), ("The reviewer rejects the proposal.", "The system records the rejection reason and allows a new proposal or revision.")),
        (("The architecture agents fail or return incomplete output.", "The runtime stores the failure evidence and prevents approval of an incomplete proposal."), ("The approval service is unavailable.", "The system preserves the pending decision and shows that approval has not been completed.")),
        "The project contains an architecture decision record with proposal, critique, rationale, reviewer, timestamp, and approval status.",
        (
            Requirement("ARCH-01", "The system shall generate at least one architecture proposal from the reviewed project requirements."),
            Requirement("ARCH-02", "The system shall represent components, APIs, data stores, technology choices, and communication patterns in a structured proposal."),
            Requirement("ARCH-03", "The Critic Agent shall evaluate security, scalability, complexity, missing components, dependencies, and requirement coverage."),
            Requirement("ARCH-04", "The system shall allow comparison of architecture alternatives and record trade-offs."),
            Requirement("ARCH-05", "The system shall store an architecture decision record with rationale and review history."),
            Requirement("ARCH-06", "The system shall require an approval state before an architecture proposal can be used for controlled code generation."),
        ),
    ),
    Feature(
        "6", "Multi-Agent Code Generation",
        "The platform uses specialized agents to generate frontend, backend, API, database, configuration, test, and documentation artifacts while the Agent Runtime controls tools, permissions, repositories, and versions.",
        "High", "Project Member; Coding Agents; Agent Runtime; LLM Provider", "Generate a coherent initial application implementation from approved requirements and architecture.", "An approved architecture decision exists, the target repository is accessible, and agent/model configuration is available.",
        (
            ("The user selects Generate Code.", "The system shows the approved requirements, architecture, target repository, agent plan, and requested artifact types."),
            ("The user confirms the generation scope.", "The Agent Runtime creates a run, assigns specialized agents, and grants only the required tools and artifact permissions."),
            ("Agents generate artifacts.", "The runtime captures prompts, outputs, file changes, dependency requests, and agent execution status."),
            ("Generation completes.", "The system creates a versioned artifact set and shows a summary/diff for review."),
            ("The user reviews the generated changes.", "The system allows approval for sandbox execution or rejection for revision."),
        ),
        (("The user requests generation without an approved architecture.", "The system blocks generation and identifies the missing approval."), ("An agent requests an unauthorized tool or file.", "The runtime denies the request and records the permission decision.")),
        (("An LLM provider is unavailable or returns an invalid artifact.", "The run is marked incomplete, the generated output is not treated as verified, and the user receives a diagnostic."), ("A repository write fails.", "The system preserves the proposed change set and reports that the repository was not updated.")),
        "The project contains a versioned generated artifact set with agent history, permissions, diffs, and an approval state for execution.",
        (
            Requirement("CODE-01", "The system shall create a generation plan from approved requirements and architecture artifacts."),
            Requirement("CODE-02", "The system shall support specialized agents for frontend, backend/API, database/configuration, tests, and documentation artifacts."),
            Requirement("CODE-03", "The Agent Runtime shall enforce tool, repository, and artifact permissions for each agent."),
            Requirement("CODE-04", "The system shall preserve generated files, diffs, dependencies, prompts or references, and execution status as a versioned artifact set."),
            Requirement("CODE-05", "The system shall allow a user to review and approve or reject a generated change set before sandbox execution."),
            Requirement("CODE-06", "The system shall not label generated code as verified until it passes the verification workflow."),
        ),
    ),
    Feature(
        "7", "Docker Sandbox Execution",
        "The Docker Sandbox builds and runs generated applications in a controlled environment and collects process output, exit status, build results, and startup results.",
        "High", "Project Member; Agent Runtime; Docker Execution Service", "Execute generated software safely enough to produce repeatable FYP-I build and runtime evidence.", "A generated artifact set is approved for execution and Docker is available.",
        (
            ("The user selects Run in Sandbox.", "The system validates the artifact version, execution policy, image configuration, and sandbox limits."),
            ("The runtime starts the sandbox.", "The system builds or loads the configured Docker image with restricted permissions and mounted project artifacts."),
            ("The container builds and starts the application.", "The system captures stdout, stderr, exit codes, health/startup output, duration, and resource-limit events."),
            ("The execution completes or stops.", "The system stores the run evidence and returns a success, failure, timeout, or blocked status."),
        ),
        (("Docker is not available.", "The system displays a setup error and does not attempt host execution."), ("A build fails.", "The system preserves logs and marks the run failed for diagnosis or future repair.")),
        (("The container exceeds a configured limit.", "The runtime stops the container, records the limit event, and marks the run as stopped by policy."), ("The sandbox service crashes.", "The system records an infrastructure exception and prevents the result from being treated as successful.")),
        "A controlled execution record exists with image/configuration details, logs, exit status, duration, and resource/policy events.",
        (
            Requirement("SBOX-01", "The system shall execute approved generated artifacts inside a Docker sandbox."),
            Requirement("SBOX-02", "The sandbox shall apply configured restrictions for permissions, time, memory, storage, and process behavior where supported."),
            Requirement("SBOX-03", "The system shall capture build output, application startup output, standard output, standard error, exit code, and duration."),
            Requirement("SBOX-04", "The system shall prevent direct host execution when a sandbox execution is required."),
            Requirement("SBOX-05", "The system shall record timeout, resource-limit, policy-block, and infrastructure-failure events."),
            Requirement("SBOX-06", "The system shall associate each sandbox run with the exact generated artifact version and project."),
        ),
    ),
    Feature(
        "8", "Initial Evidence-Based Verification",
        "The Initial Verification System evaluates generated software using available build, runtime, test, API, static-analysis, schema, and health evidence and reports whether the FYP-I verification conditions are satisfied.",
        "High", "Project Member; Verification Service; Pytest; Playwright; API Test Runner; Static Analyzer", "Determine whether generated software is supported by executable verification evidence.", "A sandbox run exists and the project has configured verification checks or test commands.",
        (
            ("The user opens the Verification view.", "The system lists available checks for the selected project and generated artifact version."),
            ("The user starts verification.", "The runtime schedules the selected checks and displays a running state."),
            ("Checks execute.", "The system runs configured tests and analyses, captures commands, outputs, durations, and exit statuses."),
            ("Verification completes.", "The system classifies each check as passed, failed, skipped, or unavailable and attaches the result to the relevant artifacts."),
            ("The user reviews the result.", "The system shows evidence, unresolved failures, and whether the artifact can proceed to the next FYP-I review stage."),
        ),
        (("No verification checks are configured.", "The system reports that verification is incomplete and does not claim a verified result."), ("A nonblocking check is skipped by policy.", "The system records the skip reason and displays it in the evidence summary.")),
        (("A test runner or static analyzer is unavailable.", "The system records the unavailable dependency and marks the corresponding evidence as unavailable."), ("A check produces malformed output.", "The system stores the raw output, marks parsing as failed, and requires human review.")),
        "The project has a verification summary with raw evidence, parsed results, linked artifacts, and an overall FYP-I verification state.",
        (
            Requirement("VER-01", "The system shall list configured verification checks for a selected generated artifact version."),
            Requirement("VER-02", "The system shall execute available build, runtime, Pytest, Playwright, API, static-analysis, schema, and health checks as configured."),
            Requirement("VER-03", "The system shall capture the command, output, duration, exit status, and environment reference for each check."),
            Requirement("VER-04", "The system shall classify each check as passed, failed, skipped, or unavailable with a reason."),
            Requirement("VER-05", "The system shall associate verification evidence with the generated artifact and relevant traceability links."),
            Requirement("VER-06", "The system shall prevent an incomplete or failed verification run from being represented as a successful verification."),
        ),
    ),
    Feature(
        "9", "Project Workspace and Artifact Management",
        "The platform provides a workspace for projects, artifacts, agent runs, approvals, versions, execution records, and verification evidence so FYP-I work remains organized and auditable.",
        "High", "Project Member; Advisor/Reviewer; Agent Runtime", "Provide one controlled place to review the current project state and its engineering history.", "The user is authenticated and has access to the workspace.",
        (
            ("The user creates or opens a project.", "The system displays the project dashboard with status, scope, recent runs, risks, approvals, and unresolved issues."),
            ("The user navigates to an artifact category.", "The system lists requirements, stories, architecture, code, tests, documentation, and evidence with IDs and versions."),
            ("The user opens an artifact.", "The system shows content, relationships, history, owner, review state, and linked agent/execution records."),
            ("The user or runtime creates a new version.", "The system records the version, timestamp, source action, and relationship changes."),
        ),
        (("The user lacks permission for a project.", "The system denies access and does not reveal project data."), ("An artifact is archived.", "The system hides it from active views while retaining its history and links.")),
        (("The storage service is unavailable.", "The system shows a recoverable error and prevents misleading confirmation of unsaved changes."), ("Two updates conflict.", "The system identifies the version conflict and asks the user or runtime to reconcile it.")),
        "The project has an organized, versioned workspace with auditable artifacts, runs, approvals, and evidence.",
        (
            Requirement("WORK-01", "The system shall allow an authenticated project member to create and open a project workspace."),
            Requirement("WORK-02", "The system shall display project status, recent agent runs, risks, approvals, unresolved issues, and verification state."),
            Requirement("WORK-03", "The system shall store and retrieve versioned requirements, user stories, architecture, code, tests, documentation, and evidence artifacts."),
            Requirement("WORK-04", "The system shall associate artifacts with their creator, timestamp, version, and source workflow action."),
            Requirement("WORK-05", "The system shall enforce project-level access control for workspace data and actions."),
            Requirement("WORK-06", "The system shall preserve an audit history for approvals, permission decisions, execution runs, and artifact changes."),
        ),
    ),
)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    srs_path = add_srs_document(FEATURES)
    use_cases_path = add_use_case_document(FEATURES)
    print(srs_path)
    print(use_cases_path)
